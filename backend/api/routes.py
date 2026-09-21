import asyncio
import base64
import io
import json
import logging
import tempfile
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import cv2
import numpy as np
import soundfile as sf
from fastapi import APIRouter, UploadFile, File, HTTPException, Request, WebSocket, WebSocketDisconnect, Depends, Response
from fastapi.responses import StreamingResponse, FileResponse

from api.schemas import (
    JobStatus,
    JobResult,
    ExportRequest,
    ExportSegment,
    LoginRequest,
    AuthUserOut,
    HistoryItem,
)
from core.auth import AuthUser, get_auth_store, require_user
from core import history as history_store
from core.face_recognizer import get_face_recognizer
from tasks.manager import JobManager, run_processing_job

logger = logging.getLogger(__name__)

router = APIRouter()

BASE_DIR = Path(__file__).parent.parent
UPLOADS_DIR = BASE_DIR / "storage" / "uploads"
JOBS_DIR = BASE_DIR / "storage" / "jobs"
STORAGE_DIR = BASE_DIR / "storage"

job_manager = JobManager(JOBS_DIR)
auth_store = get_auth_store(STORAGE_DIR)

# 线程池用于模型推理
thread_pool = ThreadPoolExecutor(max_workers=4)

# 实时识别音频配置
AUDIO_WINDOW_SECONDS = 3      # 音频窗口时长
AUDIO_SLIDE_SECONDS = 1.5     # 滑动步长
AUDIO_SAMPLE_RATE = 16000     # 采样率
SILENCE_RMS_THRESHOLD = 0.03  # 静音检测阈值（RMS 能量），低于此值视为安静

ALLOWED_EXTENSIONS = {".mp4", ".avi", ".mov", ".mkv", ".flv", ".wmv"}


def _check_job_access(job_id: str, user: AuthUser) -> None:
    try:
        job_manager.assert_owner(job_id, user.user_id)
    except KeyError:
        raise HTTPException(404, "Job not found")
    except PermissionError:
        raise HTTPException(403, "无权访问该识别记录")


# ==================== 鉴权（仅登录，不提供注册） ====================


def _client_ip(request: Request) -> str:
    """优先取反向代理传递的真实客户端 IP。"""
    xff = request.headers.get("x-forwarded-for")
    if xff:
        return xff.split(",")[0].strip() or "unknown"
    if request.client and request.client.host:
        return request.client.host
    return "unknown"


@router.post("/auth/login", response_model=AuthUserOut)
async def auth_login(payload: LoginRequest, request: Request, response: Response):
    user = auth_store.login(payload.username, payload.password, client_ip=_client_ip(request))
    token = auth_store.create_token(user)
    auth_store.set_auth_cookie(response, token)
    return AuthUserOut(user_id=user.user_id, username=user.username)


@router.post("/auth/logout")
async def auth_logout(response: Response):
    auth_store.clear_auth_cookie(response)
    return {"ok": True}


@router.get("/auth/me", response_model=AuthUserOut)
async def auth_me(user: AuthUser = Depends(require_user)):
    return AuthUserOut(user_id=user.user_id, username=user.username)


# ==================== 识别历史 ====================


@router.get("/history", response_model=list[HistoryItem])
async def get_history(user: AuthUser = Depends(require_user)):
    return history_store.list_history(JOBS_DIR, user.user_id)


@router.get("/history/{job_id}/result", response_model=JobResult)
async def history_result(job_id: str, user: AuthUser = Depends(require_user)):
    _check_job_access(job_id, user)
    result = job_manager.get_result(job_id)
    if result is None:
        status = job_manager.get_status(job_id)
        if status is not None and status.status == "failed":
            raise HTTPException(500, f"Job failed: {status.message}")
        raise HTTPException(404, "识别结果不存在")
    return result


# ==================== 上传与任务 ====================


@router.post("/upload", response_model=JobStatus)
async def upload_video(file: UploadFile = File(...), user: AuthUser = Depends(require_user)):
    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, f"Unsupported file format: {ext}. Allowed: {', '.join(ALLOWED_EXTENSIONS)}")

    # Save uploaded file
    UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    save_path = UPLOADS_DIR / f"{user.user_id}_{file.filename}"
    content = await file.read()
    save_path.write_bytes(content)

    # Create job
    job_id = job_manager.create_job(str(save_path), file.filename, user_id=user.user_id)

    # Start background processing
    thread = threading.Thread(
        target=run_processing_job,
        args=(job_manager, job_id, str(save_path), file.filename, JOBS_DIR),
        daemon=True,
    )
    thread.start()

    return job_manager.get_status(job_id)


@router.get("/jobs/{job_id}/progress")
async def job_progress(job_id: str, request: Request, user: AuthUser = Depends(require_user)):
    _check_job_access(job_id, user)
    status = job_manager.get_status(job_id)
    snapshot = job_manager.get_sse_snapshot(job_id)
    if status is None and snapshot is None:
        raise HTTPException(404, "Job not found")

    # 已结束：单次下发快照（含完整/部分片段）
    if status is not None and status.status in ("done", "failed"):
        payload = snapshot or status.model_dump()

        async def single_event():
            yield f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"
        return StreamingResponse(single_event(), media_type="text/event-stream")

    q = job_manager.subscribe_sse(job_id)

    async def event_generator():
        try:
            # 订阅后先推当前进度与已识别片段，刷新页面不丢结果
            if snapshot is not None:
                yield f"data: {json.dumps(snapshot, ensure_ascii=False)}\n\n"
                if snapshot.get("status") in ("done", "failed"):
                    return
            while True:
                if await request.is_disconnected():
                    break
                try:
                    data = await asyncio.to_thread(q.get, timeout=30)
                    yield f"data: {data}\n\n"
                    parsed = json.loads(data)
                    if parsed.get("status") in ("done", "failed") or parsed.get("event") in ("done", "failed"):
                        break
                except Exception:
                    yield ": keepalive\n\n"
        finally:
            job_manager.cleanup_sse(job_id, q)

    return StreamingResponse(event_generator(), media_type="text/event-stream")


@router.get("/jobs/{job_id}/result", response_model=JobResult)
async def job_result(job_id: str, user: AuthUser = Depends(require_user)):
    _check_job_access(job_id, user)
    result = job_manager.get_result(job_id)
    if result is None:
        status = job_manager.get_status(job_id)
        if status is None:
            raise HTTPException(404, "Job not found")
        if status.status == "failed":
            raise HTTPException(500, f"Job failed: {status.message}")
        raise HTTPException(202, "Job still processing")
    # 处理中返回 partial（complete=false），前端可持续合并片段
    return result


@router.get("/jobs/{job_id}/original-video")
async def job_original_video(job_id: str, user: AuthUser = Depends(require_user)):
    _check_job_access(job_id, user)
    video_path = job_manager.get_video_path(job_id)
    if video_path is None or not Path(video_path).is_file():
        raise HTTPException(404, "Original video not found")
    return FileResponse(video_path, media_type="video/mp4")


@router.get("/jobs/{job_id}/video/{segment_index}")
async def job_video_segment(job_id: str, segment_index: int, user: AuthUser = Depends(require_user)):
    _check_job_access(job_id, user)
    segments_dir = JOBS_DIR / job_id / "segments"
    if not segments_dir.exists():
        raise HTTPException(404, "Segments not found")

    matches = list(segments_dir.glob(f"*_seg{segment_index:04d}.mp4"))
    if not matches:
        raise HTTPException(404, f"Segment {segment_index} not found")

    return FileResponse(matches[0], media_type="video/mp4")


def _format_time(seconds: float) -> str:
    """将秒数格式化为 HH:MM:SS 格式"""
    total_seconds = int(seconds)
    hours = total_seconds // 3600
    minutes = (total_seconds % 3600) // 60
    secs = total_seconds % 60
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _build_excel_bytes(rows: list[dict], headers: list[str], widths: list[int]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "情感识别结果"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for row_idx, row in enumerate(rows, 2):
        for col, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border

    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    if rows:
        last_col = openpyxl.utils.get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/jobs/{job_id}/export")
async def job_export_excel(job_id: str, user: AuthUser = Depends(require_user)):
    _check_job_access(job_id, user)
    result = job_manager.get_result(job_id)
    if result is None:
        status = job_manager.get_status(job_id)
        if status is None:
            raise HTTPException(404, "Job not found")
        if status.status == "failed":
            raise HTTPException(500, f"Job failed: {status.message}")
        raise HTTPException(202, "Job still processing")

    # 暂时只导出通用 emotion2vec 情感，隐藏教师情感列
    headers = [
        "#", "开始时间", "结束时间", "时长", "文本",
        "emotion2vec标签", "emotion2vec标签名", "置信度",
    ]
    rows = []
    for seg in result.segments:
        rows.append([
            seg.index + 1,
            _format_time(seg.start_time),
            _format_time(seg.end_time),
            _format_time(seg.duration),
            seg.text,
            seg.emotion2vec_label,
            seg.emotion2vec_label_name,
            round(seg.confidence, 4),
        ])

    widths = [5, 10, 10, 10, 40, 14, 16, 8]
    data = _build_excel_bytes(rows, headers, widths)

    filename = f"{result.video_name}_emotions.xlsx"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/jobs/{job_id}/export")
async def job_export_excel_edited(job_id: str, payload: ExportRequest, user: AuthUser = Depends(require_user)):
    """使用前端编辑后的情感标签与识别文本导出 Excel（当前只导出通用 emotion2vec）"""
    _check_job_access(job_id, user)
    video_name = payload.video_name or job_id
    # 暂时强制通用情感导出，忽略教师情感
    headers = [
        "#", "开始时间", "结束时间", "时长",
        "识别文本", "最终文本", "文本是否修改",
        "识别-情感标签", "识别-情感标签名",
        "最终-情感标签", "最终-情感标签名",
        "置信度", "情感是否修改",
    ]
    widths = [5, 10, 10, 10, 32, 32, 10, 14, 16, 14, 16, 8, 10]

    rows = []
    for seg in payload.segments:
        original_text = seg.original_text if seg.original_text else seg.text
        rows.append([
            seg.index + 1,
            _format_time(seg.start_time),
            _format_time(seg.end_time),
            _format_time(seg.duration),
            original_text,
            seg.text,
            "是" if seg.text_edited else "否",
            seg.original_emotion2vec_label,
            seg.original_emotion2vec_label_name or "",
            seg.emotion2vec_label,
            seg.emotion2vec_label_name or "",
            round(seg.confidence, 4),
            "是" if seg.edited else "否",
        ])

    data = _build_excel_bytes(rows, headers, widths)
    safe_name = (video_name or "result").rsplit(".", 1)[0]
    filename = f"{safe_name}_emotions.xlsx"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ==================== 实时情感识别 WebSocket ====================


def _process_video_frame(base64_data: str) -> dict:
    """处理视频帧：base64 解码 → 人脸检测 → 表情识别"""
    face_recognizer = get_face_recognizer()
    # 通用情感批处理场景下启动时不预加载人脸；实时用到时再载入
    if not face_recognizer.is_loaded:
        try:
            face_recognizer.load()
        except Exception as e:
            logger.error(f"人脸模型按需加载失败: {e}")
            return {"type": "face_emotion", "face_detected": False}

    if not face_recognizer.is_loaded:
        return {"type": "face_emotion", "face_detected": False}

    try:
        # base64 解码为图像
        img_bytes = base64.b64decode(base64_data)
        nparr = np.frombuffer(img_bytes, np.uint8)
        frame = cv2.imdecode(nparr, cv2.IMREAD_COLOR)

        if frame is None:
            return {"type": "face_emotion", "face_detected": False}

        # 推理
        result = face_recognizer.predict(frame)
        result["type"] = "face_emotion"
        return result
    except Exception as e:
        logger.error(f"视频帧处理失败: {e}")
        return {"type": "face_emotion", "face_detected": False}


def _process_audio_buffer(audio_buffer: np.ndarray) -> dict:
    """处理音频缓冲：PCM → WAV → emotion2vec 推理"""
    from core.recognizer import _get_predictor
    from core.face_recognizer import map_emotion2vec_to_7class

    try:
        # 静音检测：计算 RMS 能量
        rms = float(np.sqrt(np.mean(audio_buffer ** 2)))
        if rms < SILENCE_RMS_THRESHOLD:
            return {"type": "voice_emotion", "silence": True}

        # 保存为临时 WAV 文件
        with tempfile.NamedTemporaryFile(suffix=".wav", delete=False) as tmp:
            tmp_path = tmp.name
            sf.write(tmp_path, audio_buffer, AUDIO_SAMPLE_RATE)

        # emotion2vec 推理
        predictor = _get_predictor()
        raw_result = predictor.predict(tmp_path)

        # 清理临时文件
        import os
        try:
            os.remove(tmp_path)
        except OSError:
            pass

        # 映射到 7 类
        result = map_emotion2vec_to_7class(raw_result)
        result["type"] = "voice_emotion"
        return result
    except Exception as e:
        logger.error(f"音频处理失败: {e}")
        return {"type": "voice_emotion", "silence": True}


@router.websocket("/ws/realtime")
async def websocket_realtime(websocket: WebSocket):
    """实时情感识别 WebSocket 端点"""
    await websocket.accept()

    # 创建实时识别 Job
    job_id = uuid.uuid4().hex[:12]
    job_dir = JOBS_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)

    await websocket.send_json({"type": "connected", "job_id": job_id})
    logger.info(f"实时识别 WebSocket 已连接, job_id={job_id}")

    # 音频缓冲区
    audio_buffer = np.array([], dtype=np.float32)

    # 结果累积
    results = []
    start_time = time.time()

    # 跟踪并发任务
    pending_tasks = set()

    try:
        while True:
            raw = await websocket.receive_text()
            message = json.loads(raw)
            msg_type = message.get("type")

            if msg_type == "video_frame":
                # 在线程池中处理视频帧
                loop = asyncio.get_event_loop()
                task = loop.run_in_executor(
                    thread_pool,
                    _process_video_frame,
                    message["data"],
                )
                pending_tasks.add(task)
                task.add_done_callback(pending_tasks.discard)

                # 异步等待结果并推送
                async def send_face_result(t):
                    try:
                        result = await t
                        results.append({
                            "timestamp": time.time() - start_time,
                            "source": "face",
                            **result,
                        })
                        await websocket.send_json(result)
                    except Exception as e:
                        logger.error(f"发送面部结果失败: {e}")

                asyncio.create_task(send_face_result(task))

            elif msg_type == "audio_chunk":
                # base64 解码音频数据（PCM float32）
                try:
                    audio_bytes = base64.b64decode(message["data"])
                    pcm_data = np.frombuffer(audio_bytes, dtype=np.float32)
                    audio_buffer = np.concatenate([audio_buffer, pcm_data])
                except Exception as e:
                    logger.error(f"音频解码失败: {e}")
                    continue

                # 当缓冲区达到窗口大小时触发推理
                window_samples = AUDIO_WINDOW_SECONDS * AUDIO_SAMPLE_RATE
                if len(audio_buffer) >= window_samples:
                    # 取出一个窗口
                    window = audio_buffer[:window_samples].copy()

                    # 滑动窗口：保留后半部分
                    slide_samples = int(AUDIO_SLIDE_SECONDS * AUDIO_SAMPLE_RATE)
                    audio_buffer = audio_buffer[slide_samples:]

                    # 在线程池中处理音频
                    loop = asyncio.get_event_loop()
                    task = loop.run_in_executor(
                        thread_pool,
                        _process_audio_buffer,
                        window,
                    )
                    pending_tasks.add(task)
                    task.add_done_callback(pending_tasks.discard)

                    async def send_voice_result(t):
                        try:
                            result = await t
                            results.append({
                                "timestamp": time.time() - start_time,
                                "source": "voice",
                                **result,
                            })
                            await websocket.send_json(result)
                        except Exception as e:
                            logger.error(f"发送语音结果失败: {e}")

                    asyncio.create_task(send_voice_result(task))

    except WebSocketDisconnect:
        logger.info(f"实时识别 WebSocket 断开, job_id={job_id}")
    except Exception as e:
        logger.error(f"WebSocket 错误: {e}")
    finally:
        # 等待所有待处理任务完成
        if pending_tasks:
            await asyncio.gather(*pending_tasks, return_exceptions=True)

        # 保存结果
        duration = time.time() - start_time
        realtime_result = {
            "job_id": job_id,
            "mode": "realtime",
            "duration_seconds": round(duration, 2),
            "start_time": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(start_time)),
            "end_time": time.strftime("%Y-%m-%dT%H:%M:%SZ"),
            "timeline": [
                {
                    "timestamp": r.get("timestamp", 0),
                    "source": r.get("source", "face"),
                    "emotion": r.get("emotion", "neutral"),
                    "confidence": r.get("confidence", 0),
                    "probabilities": r.get("probabilities", {}),
                }
                for r in results
            ],
            "summary": _compute_summary(results),
        }

        result_path = job_dir / "result.json"
        result_path.write_text(json.dumps(realtime_result, ensure_ascii=False, indent=2), encoding="utf-8")
        logger.info(f"实时识别结果已保存: {result_path}")


def _compute_summary(results: list) -> dict:
    """计算情感分布统计"""
    emotions = ["angry", "disgust", "fear", "happy", "neutral", "sad", "surprise"]
    summary = {
        "face": {e: 0 for e in emotions},
        "voice": {e: 0 for e in emotions},
    }
    for r in results:
        source = r.get("source", "face")
        emotion = r.get("emotion", "neutral")
        if source in summary and emotion in summary[source]:
            summary[source][emotion] += 1
    return summary
