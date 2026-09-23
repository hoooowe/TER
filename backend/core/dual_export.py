"""导出「教学情绪-行为」标准化双维编码表（Excel）。"""

from __future__ import annotations

import io
from typing import Any, Iterable

from core.coding_framework import (
    BEHAVIOR_CODES,
    BEHAVIOR_ORDER,
    EMOTION_CODES,
    EMOTION_ORDER,
    build_transition_matrices,
    compute_process_metrics,
)


def _format_time(seconds: float) -> str:
    total_seconds = int(seconds)
    hours = total_seconds // 3600
    minutes = (total_seconds % 3600) // 60
    secs = total_seconds % 60
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _style_header(ws, headers: list[str], widths: list[int] | None = None):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _write_matrix_sheet(wb, title: str, matrix: dict[str, dict[str, int]], row_order: list[str], col_order: list[str]):
    ws = wb.create_sheet(title)
    ws.cell(row=1, column=1, value="前一状态 \\ 后一状态")
    for j, c in enumerate(col_order, 2):
        ws.cell(row=1, column=j, value=c)
    for i, r in enumerate(row_order, 2):
        ws.cell(row=i, column=1, value=r)
        for j, c in enumerate(col_order, 2):
            ws.cell(row=i, column=j, value=matrix.get(r, {}).get(c, 0))


def build_dual_excel_bytes(
    units: Iterable[Any],
    *,
    video_name: str = "",
    include_matrices: bool = True,
) -> bytes:
    """
    生成标准化双维编码表 Excel。

    Sheet1 双维编码表：# / 开始 / 结束 / 时长 / 文本 / 行为码 / 行为名称 / 行为维度 /
                    行为依据 / 情绪码 / 情绪名称 / 情绪维度 / 外显线索 / 置信度 / 待复核
    Sheet2 编码手册
    Sheet3 过程指标
    Sheet4-6 矩阵（可选）
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    unit_list = list(units)
    wb = openpyxl.Workbook()

    # ---------- Sheet1 双维编码表 ----------
    ws = wb.active
    ws.title = "双维编码表"
    headers = [
        "#", "开始时间", "结束时间", "时长", "文本",
        "行为码", "行为名称", "行为维度", "行为判定依据",
        "情绪码", "情绪名称", "情绪维度", "外显情绪线索",
        "置信度", "待人工复核",
    ]
    widths = [5, 10, 10, 8, 36, 8, 12, 12, 28, 8, 12, 12, 28, 8, 10]
    _style_header(ws, headers, widths)

    review_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for row_idx, u in enumerate(unit_list, 2):
        values = [
            getattr(u, "index", row_idx - 2) + 1 if hasattr(u, "index") else row_idx - 1,
            _format_time(getattr(u, "start_time", 0)),
            _format_time(getattr(u, "end_time", 0)),
            _format_time(getattr(u, "duration", 0) or (getattr(u, "end_time", 0) - getattr(u, "start_time", 0))),
            getattr(u, "text", "") or "",
            getattr(u, "behavior_code", ""),
            getattr(u, "behavior_name", ""),
            getattr(u, "behavior_dim", ""),
            getattr(u, "behavior_reason", ""),
            getattr(u, "emotion_code", ""),
            getattr(u, "emotion_name", ""),
            getattr(u, "emotion_dim", ""),
            getattr(u, "emotion_evidence", ""),
            round(float(getattr(u, "confidence", 0) or 0), 4),
            "是" if getattr(u, "needs_review", False) else "否",
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if getattr(u, "needs_review", False):
                cell.fill = review_fill

    if unit_list:
        last_col = openpyxl.utils.get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{len(unit_list) + 1}"

    # ---------- Sheet2 编码手册 ----------
    ws2 = wb.create_sheet("编码手册")
    _style_header(ws2, ["维度", "编码", "名称", "描述"], [12, 8, 14, 60])
    r = 2
    for e in EMOTION_CODES:
        ws2.cell(row=r, column=1, value="情绪-" + e["dim"])
        ws2.cell(row=r, column=2, value=e["code"])
        ws2.cell(row=r, column=3, value=e["name"])
        ws2.cell(row=r, column=4, value=e["description"])
        r += 1
    for b in BEHAVIOR_CODES:
        ws2.cell(row=r, column=1, value="行为-" + b["dim"])
        ws2.cell(row=r, column=2, value=b["code"])
        ws2.cell(row=r, column=3, value=b["name"])
        ws2.cell(row=r, column=4, value=b["description"])
        r += 1

    # ---------- Sheet3 过程指标 ----------
    metrics = compute_process_metrics(unit_list)
    ws3 = wb.create_sheet("过程指标")
    _style_header(ws3, ["指标缩写", "中文名称", "数值", "含义"], [12, 18, 10, 50])
    metric_rows = [
        ("KPBR", "知识呈现行为占比", metrics.get("KPBR", 0), "知识讲解、举例说明、归纳总结等占全部教学行为的比例"),
        ("IIBR", "教学互动行为占比", metrics.get("IIBR", 0), "提问、反馈、评价等互动行为占全部教学行为的比例"),
        ("IOBR", "教学操作行为占比", metrics.get("IOBR", 0), "板书、课件操作、教学演示等占全部教学行为的比例"),
        ("COBR", "课堂组织行为占比", metrics.get("COBR", 0), "活动组织、活动说明、环节过渡等占全部教学行为的比例"),
        ("TEOR", "失误行为发生率", metrics.get("TEOR", 0), "授课失误行为占全部教学行为事件的比例"),
        ("PER", "积极情绪出现率", metrics.get("PER", 0), "积极情绪（E1-E3）占全部情绪编码事件的比例"),
        ("NER", "中性情绪出现率", metrics.get("NER", 0), "中性情绪（U1）占全部情绪编码事件的比例"),
        ("NGR", "消极情绪出现率", metrics.get("NGR", 0), "消极情绪（N1-N2）占全部情绪编码事件的比例"),
    ]
    for i, row in enumerate(metric_rows, 2):
        for c, v in enumerate(row, 1):
            ws3.cell(row=i, column=c, value=v)

    if include_matrices:
        mats = build_transition_matrices(unit_list)
        _write_matrix_sheet(
            wb, "行为连接矩阵",
            mats["behavior_transition"], BEHAVIOR_ORDER, BEHAVIOR_ORDER,
        )
        _write_matrix_sheet(
            wb, "情绪转换矩阵",
            mats["emotion_transition"], EMOTION_ORDER, EMOTION_ORDER,
        )
        # 共现：行=情绪，列=行为
        _write_matrix_sheet(
            wb, "行为情绪共现矩阵",
            mats["emotion_behavior_cooccurrence"], EMOTION_ORDER, BEHAVIOR_ORDER,
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
